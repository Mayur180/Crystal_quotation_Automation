from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# ============================================================
# CRYSTAL / ARMATECH ASSOCIATES
# WORD DOCUMENT COLORS
# ============================================================

RED = RGBColor(192, 0, 0)
DARK_GREY = RGBColor(70, 70, 70)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)


class WordStyleService:
    """
    Reusable Word document styling service for
    ARMATECH ASSOCIATES quotations.
    """

    def __init__(self):
        self.font_name = "Arial"

    # ========================================================
    # DOCUMENT CREATION
    # ========================================================

    def create_document(self) -> Document:
        """
        Create a new A4 Word document with
        ARMATECH standard formatting.
        """

        document = Document()

        self.setup_page(document)
        self.setup_styles(document)
        self.setup_header_footer(document)

        return document

    # ========================================================
    # PAGE SETUP
    # ========================================================

    def setup_page(self, document: Document) -> None:
        """
        Configure A4 page size and margins.
        """

        for section in document.sections:

            section.page_width = Inches(8.27)
            section.page_height = Inches(11.69)

            section.top_margin = Inches(0.65)
            section.bottom_margin = Inches(0.60)
            section.left_margin = Inches(0.65)
            section.right_margin = Inches(0.65)

            section.header_distance = Inches(0.25)
            section.footer_distance = Inches(0.25)

    # ========================================================
    # WORD STYLES
    # ========================================================

    def setup_styles(self, document: Document) -> None:
        """
        Configure standard Word document styles.
        """

        # ----------------------------------------------------
        # NORMAL
        # ----------------------------------------------------

        normal = document.styles["Normal"]

        normal.font.name = self.font_name
        normal.font.size = Pt(9)

        normal._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            self.font_name
        )

        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.line_spacing = 1.0

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title = document.styles["Title"]

        title.font.name = self.font_name
        title.font.size = Pt(16)
        title.font.bold = True
        title.font.color.rgb = RED

        title._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            self.font_name
        )

        # ----------------------------------------------------
        # HEADING 1
        # ----------------------------------------------------

        heading1 = document.styles["Heading 1"]

        heading1.font.name = self.font_name
        heading1.font.size = Pt(12)
        heading1.font.bold = True
        heading1.font.color.rgb = RED

        heading1._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            self.font_name
        )

        # ----------------------------------------------------
        # HEADING 2
        # ----------------------------------------------------

        heading2 = document.styles["Heading 2"]

        heading2.font.name = self.font_name
        heading2.font.size = Pt(10)
        heading2.font.bold = True
        heading2.font.color.rgb = DARK_GREY

        heading2._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            self.font_name
        )

    # ========================================================
    # HEADER / FOOTER
    # ========================================================

    def setup_header_footer(self, document: Document) -> None:
        """
        Configure ARMATECH ASSOCIATES header and footer.
        """

        for section in document.sections:

            # =================================================
            # HEADER
            # =================================================

            header = section.header

            header_table = header.add_table(
                rows=1,
                cols=2,
                width=Inches(6.97)
            )

            header_table.autofit = False

            header_table.columns[0].width = Inches(1.65)
            header_table.columns[1].width = Inches(5.32)

            left_cell = header_table.cell(0, 0)
            right_cell = header_table.cell(0, 1)

            left_cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            right_cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            # -------------------------------------------------
            # LEFT SIDE - ARMATECH
            # -------------------------------------------------

            left_paragraph = left_cell.paragraphs[0]

            left_paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
            )

            logo_text = left_paragraph.add_run(
                "ARMATECH"
            )

            logo_text.bold = True
            logo_text.font.name = self.font_name
            logo_text.font.size = Pt(14)
            logo_text.font.color.rgb = RED

            # -------------------------------------------------
            # RIGHT SIDE - COMPANY NAME
            # -------------------------------------------------

            right_paragraph = right_cell.paragraphs[0]

            right_paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            company_run = right_paragraph.add_run(
                "ARMATECH ASSOCIATES"
            )

            company_run.bold = True
            company_run.font.name = self.font_name
            company_run.font.size = Pt(16)
            company_run.font.color.rgb = RED

            # -------------------------------------------------
            # ADDRESS
            # -------------------------------------------------

            address_paragraph = right_cell.add_paragraph()

            address_paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            address_run = address_paragraph.add_run(
                "W-174A, 'S' Block, M.I.D.C., "
                "Bhosari Industrial Estate, "
                "PUNE - 411 026. (INDIA)"
            )

            address_run.font.name = self.font_name
            address_run.font.size = Pt(7)
            address_run.font.color.rgb = DARK_GREY

            # -------------------------------------------------
            # CONTACT
            # -------------------------------------------------

            contact_paragraph = right_cell.add_paragraph()

            contact_paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            contact_run = contact_paragraph.add_run(
                "+91 7020991684 / +91 9699820070"
                "  |  armatech@gmail.com"
                "  |  www.armatechassociates.com"
            )

            contact_run.font.name = self.font_name
            contact_run.font.size = Pt(7)
            contact_run.font.color.rgb = DARK_GREY

            # -------------------------------------------------
            # REMOVE HEADER TABLE BORDERS
            # -------------------------------------------------

            self.remove_table_borders(header_table)

            # =================================================
            # FOOTER
            # =================================================

            footer = section.footer

            footer_paragraph = footer.paragraphs[0]

            footer_paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            footer_run = footer_paragraph.add_run(
                "ARMATECH ASSOCIATES  |  "
                "Precision Measurement Technology"
            )

            footer_run.font.name = self.font_name
            footer_run.font.size = Pt(7)
            footer_run.font.color.rgb = DARK_GREY

            separator_run = footer_paragraph.add_run(
                "  |  "
            )

            separator_run.font.name = self.font_name
            separator_run.font.size = Pt(7)
            separator_run.font.color.rgb = DARK_GREY

            page_run = footer_paragraph.add_run(
                "Page "
            )

            page_run.font.name = self.font_name
            page_run.font.size = Pt(7)
            page_run.font.color.rgb = DARK_GREY

            # -------------------------------------------------
            # PAGE NUMBER FIELD
            # -------------------------------------------------

            field = OxmlElement("w:fldSimple")

            field.set(
                qn("w:instr"),
                "PAGE"
            )

            footer_paragraph._p.append(field)

    # ========================================================
    # REMOVE TABLE BORDERS
    # ========================================================

    def remove_table_borders(self, table) -> None:
        """
        Remove visible borders from a table.
        """

        tbl = table._tbl
        tbl_pr = tbl.tblPr

        borders = tbl_pr.first_child_found_in(
            "w:tblBorders"
        )

        if borders is None:

            borders = OxmlElement(
                "w:tblBorders"
            )

            tbl_pr.append(borders)

        for edge in (
            "top",
            "left",
            "bottom",
            "right",
            "insideH",
            "insideV",
        ):

            tag = "w:" + edge

            element = borders.find(
                qn(tag)
            )

            if element is None:

                element = OxmlElement(tag)
                borders.append(element)

            element.set(
                qn("w:val"),
                "nil"
            )

    # ========================================================
    # CELL SHADING
    # ========================================================

    def set_cell_shading(
        self,
        cell,
        fill
    ) -> None:
        """
        Set the background/shading color
        of a Word table cell.

        Accepted values:

            "C00000"
            "#C00000"
            RGBColor(192, 0, 0)

        The Word XML layer requires the final
        value to be a plain six-character HEX string.
        """

        if fill is None:
            return

        # ----------------------------------------------------
        # Convert python-docx RGBColor to string
        # ----------------------------------------------------

        if isinstance(fill, RGBColor):

            fill = str(fill)

        elif hasattr(fill, "rgb"):

            fill = str(fill.rgb)

        else:

            fill = str(fill)

        # ----------------------------------------------------
        # Normalize HEX value
        # ----------------------------------------------------

        fill = (
            fill
            .replace("#", "")
            .strip()
            .upper()
        )

        # ----------------------------------------------------
        # Validate HEX value
        # ----------------------------------------------------

        if len(fill) != 6:

            raise ValueError(
                f"Invalid cell shading color: {fill!r}. "
                "Expected a 6-character HEX color."
            )

        # ----------------------------------------------------
        # Get table cell properties
        # ----------------------------------------------------

        tc_pr = cell._tc.get_or_add_tcPr()

        # ----------------------------------------------------
        # Remove existing shading
        # ----------------------------------------------------

        existing_shd = tc_pr.find(
            qn("w:shd")
        )

        if existing_shd is not None:

            tc_pr.remove(existing_shd)

        # ----------------------------------------------------
        # Create new shading
        # ----------------------------------------------------

        shd = OxmlElement("w:shd")

        shd.set(
            qn("w:val"),
            "clear"
        )

        shd.set(
            qn("w:color"),
            "auto"
        )

        shd.set(
            qn("w:fill"),
            fill
        )

        tc_pr.append(shd)

    # ========================================================
    # CELL TEXT
    # ========================================================

    def set_cell_text(
        self,
        cell,
        text: str,
        bold: bool = False,
        size: int = 9,
        color: RGBColor = BLACK,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    ) -> None:
        """
        Write formatted text into a table cell.
        """

        cell.text = ""

        paragraph = cell.paragraphs[0]

        paragraph.alignment = alignment

        run = paragraph.add_run(
            str(text)
        )

        run.font.name = self.font_name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

        cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

    # ========================================================
    # SECTION HEADING
    # ========================================================

    def add_section_heading(
        self,
        document: Document,
        text: str,
    ):
        """
        Add a standard red section heading.
        """

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(4)

        run = paragraph.add_run(
            str(text)
        )

        run.font.name = self.font_name
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RED

        return paragraph


# ============================================================
# DIRECT TEST
# ============================================================

if __name__ == "__main__":

    print("=" * 60)
    print("CRYSTAL WORD STYLE SERVICE TEST")
    print("=" * 60)

    service = WordStyleService()

    document = service.create_document()

    # --------------------------------------------------------
    # TEST SECTION HEADING
    # --------------------------------------------------------

    service.add_section_heading(
        document,
        "AUTOMATIC VIDEO MEASURING MACHINE"
    )

    # --------------------------------------------------------
    # TEST NORMAL TEXT
    # --------------------------------------------------------

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        "This is a Word document styling test."
    )

    run.font.name = "Arial"
    run.font.size = Pt(10)

    # --------------------------------------------------------
    # TEST TABLE
    # --------------------------------------------------------

    table = document.add_table(
        rows=2,
        cols=3
    )

    table.autofit = False

    headers = [
        "Sr. No.",
        "Product",
        "Model"
    ]

    values = [
        "1",
        "Automatic Video Measuring Machine",
        "AA-9043AC"
    ]

    # Header row
    for index, value in enumerate(headers):

        cell = table.cell(0, index)

        service.set_cell_shading(
            cell,
            RED
        )

        service.set_cell_text(
            cell,
            value,
            bold=True,
            size=9,
            color=WHITE,
            alignment=WD_ALIGN_PARAGRAPH.CENTER
        )

    # Data row
    for index, value in enumerate(values):

        cell = table.cell(1, index)

        service.set_cell_text(
            cell,
            value,
            bold=False,
            size=9,
            color=BLACK,
            alignment=WD_ALIGN_PARAGRAPH.CENTER
        )

    # --------------------------------------------------------
    # OUTPUT DIRECTORY
    # --------------------------------------------------------

    output_directory = (
        Path(__file__).resolve().parents[2]
        / "output"
        / "quotations"
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True
    )

    # --------------------------------------------------------
    # OUTPUT FILE
    # --------------------------------------------------------

    output_file = (
        output_directory
        / "word_style_test.docx"
    )

    document.save(
        output_file
    )

    print()
    print(f"Created: {output_file}")
    print()
    print("WORD STYLE SERVICE: OK")